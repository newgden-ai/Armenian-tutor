#!/usr/bin/env python3
"""Курсы «Грамматика», «История», «Обществознание», «Культура»: xlsx → SQL (повторный запуск безопасен, прогресс сохраняется).
Армянская версия: колонки «… (հայերեն)», разделы и уровни тем — армянские; внутренние поля *_uz = армянский текст.

Запуск:  python3 scripts/build_courses.py data/source data/build/courses.sql
Файлы:   data/source/grammar.xlsx, history.xlsx, civics.xlsx (имена — как ниже в SOURCES).

Что делает:
  • темы (кратко / подробно на RU-HY-EN, источник), 5 заданий на тему (без ответов → «открытые»);
  • разделы, периоды истории, хронология, справочник контактов, базовая лексика грамматики;
  • к темам грамматики — шаблоны поиска примеров в словаре (GRAMMAR_PATTERNS), их можно менять.
"""
import json
import re
import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))
from build_vocab import key_uz  # noqa: E402  ключ узбекского слова — как в словаре

SOURCES = {'grammar': 'grammar.xlsx', 'history': 'history.xlsx', 'civics': 'civics.xlsx', 'culture': 'culture.xlsx'}

# Разделы: название RU (как в xlsx) → (армянское, английское)
SECTION_TR = {
    'Фонетика': ('Հնչյունաբանություն', 'Phonetics'),
    'Морфология: имя': ('Ձևաբանություն. անուն', 'Morphology: nominals'),
    'Морфология: местоимения': ('Ձևաբանություն. դերանուններ', 'Morphology: pronouns'),
    'Морфология: глагол': ('Ձևաբանություն. բայ', 'Morphology: the verb'),
    'Синтаксис': ('Շարահյուսություն', 'Syntax'),
    'Живой язык': ('Կենդանի լեզու', 'Living language'),
    'Государство и право': ('Պետություն և իրավունք', 'State and law'),
    'Право в повседневной жизни': ('Իրավունքը առօրյա կյանքում', 'Law in everyday life'),
    'Для иностранцев в Армении': ('Օտարերկրացիների համար Հայաստանում', 'For foreigners in Armenia'),
    'Обществознание': ('Հասարակագիտություն', 'Social studies'),
    'Гостеприимство и повседневная жизнь': ('Հյուրընկալություն և առօրյա կյանք', 'Hospitality and everyday life'),
    'Семейные обряды': ('Ընտանեկան ծեսեր', 'Family rituals'),
    'Праздники': ('Տոներ', 'Holidays'),
    'Ремёсла, одежда, искусство': ('Արհեստներ, տարազ, արվեստ', 'Crafts, clothing and art'),
    # диалоги
    'Знакомство и общение': ('Ծանոթություն և շփում', 'Meeting people'),
    'Город и транспорт': ('Քաղաք և տրանսպորտ', 'City and transport'),
    'Покупки и еда': ('Գնումներ և սնունդ', 'Shopping and food'),
    'Быт и услуги': ('Կենցաղ և ծառայություններ', 'Everyday life and services'),
    'Работа и учёба': ('Աշխատանք և ուսում', 'Work and study'),
    'Здоровье': ('Առողջություն', 'Health'),
    'Праздники и гости': ('Տոներ և հյուրեր', 'Holidays and guests'),
}
SECTION_ORDER = {
    'dialogs': ['Знакомство и общение', 'Город и транспорт', 'Покупки и еда', 'Быт и услуги', 'Работа и учёба',
                'Здоровье', 'Праздники и гости'],
    'grammar': ['Фонетика', 'Морфология: имя', 'Морфология: местоимения', 'Морфология: глагол', 'Синтаксис', 'Живой язык'],
    'civics': ['Государство и право', 'Право в повседневной жизни', 'Для иностранцев в Армении', 'Обществознание'],
    'culture': ['Гостеприимство и повседневная жизнь', 'Семейные обряды', 'Праздники', 'Ремёсла, одежда, искусство'],
}

# Уровень тем (A1–B2): по сложности языка и по тому, когда тема нужна учащемуся. Не указано — DEFAULT_CEFR.
UNIT_CEFR = {
    'grammar': {6: 'A2', 11: 'A2', 12: 'A2', 15: 'A2', 19: 'A2', 13: 'B1', 20: 'B1',
                21: 'A2', 22: 'B1', 23: 'B1', 24: 'A2', 25: 'B1', 26: 'B2', 27: 'A2', 28: 'B1', 29: 'A2', 30: 'A1',
                31: 'B1', 32: 'A1', 33: 'B1'},
    'culture': {5: 'A2', 6: 'A2', 7: 'A2', 8: 'A2', 9: 'A2', 10: 'A2', 15: 'A2', 16: 'A2', 19: 'A2', 20: 'A2',
                13: 'B1', 17: 'B1', 18: 'B1'},
    'history': {8: 'A2', 21: 'A2', 22: 'A2', 1: 'B2', 2: 'B2', 4: 'B2', 7: 'B2', 11: 'B2', 14: 'B2', 15: 'B2'},
    'civics': {10: 'A2', 11: 'A2', 17: 'A2', 18: 'A2', 20: 'A2', 3: 'B2', 6: 'B2', 7: 'B2', 9: 'B2', 13: 'B2', 14: 'B2'},
}
DEFAULT_CEFR = {'grammar': 'A1', 'culture': 'A1', 'history': 'B1', 'civics': 'B1'}

# Армянские термины темы: слова в скобках в русском тексте — «лаваш (լավաշ)», «(թոնիր)».
# Те, что есть в словаре, после прочтения темы попадают в ближайшую очередь заучивания.
TERM_RE = re.compile(r'\(([^()]{2,80})\)')
CYR = re.compile(r'[А-Яа-яЁё]')
HY = re.compile(r'[ա-ևԱ-Ֆ]')


def unit_terms(*texts):
    out = set()
    for text in texts:
        for m in TERM_RE.finditer(text or ''):
            for part in re.split(r'[,;/]| или | и ', m.group(1)):
                part = part.strip(' «»"\'.')
                if not part or CYR.search(part) or not HY.search(part) or re.search(r'\d', part) or len(part) > 40:
                    continue
                k = key_uz(part)
                if k and len(k) >= 2:
                    out.add(k)
    return sorted(out)


# Названия тем грамматики на армянском и английском (в файле одно поле «Тема (RU / HY / EN)»:
# «Русское название (Հայերեն անվանում)»; армянское берётся из скобок, если его нет здесь)
GRAMMAR_TITLES_EN = {
    1: 'The Armenian alphabet',
    2: 'Pronunciation of key Armenian sounds',
    3: 'Nouns: no gender, the plural',
    4: 'The case system: 7 cases',
    5: 'Possessive pronouns and articles',
    6: 'Adjectives and degrees of comparison',
    7: 'Numerals: cardinal and ordinal',
    8: 'Personal and demonstrative pronouns',
    9: 'Interrogative pronouns',
    10: 'Present tense: -ում եմ',
    11: 'Imperfect: -ում էի',
    12: 'Simple past (aorist)',
    13: 'Present perfect and pluperfect',
    14: 'Negation (չ-) and prohibition (մի՛)',
    15: 'Imperative mood',
    16: 'SOV word order and the place of modifiers',
    17: 'Existence and possession: կա/չկա, ունեմ/չունեմ',
    18: 'Conjunctions: և, բայց, կամ, որովհետև, ուստի',
    19: 'Questions and particles (մի՞թե, արդյո՞ք)',
    20: 'Direct and indirect speech',
    21: 'Future tense and the conditional future',
    22: 'The subjunctive (optative) mood',
    23: 'Conditional sentences with եթե',
    24: 'Modal words: պետք է, կարող եմ, կարելի է, պիտի',
    25: 'Participles and converbs: −ած, −ող, −ելիս, −ելով',
    26: 'Passive and causative verbs: −վ−, −ցն−',
    27: 'Postpositions: մեջ, վրա, տակ, մոտ, հետ, համար, մասին',
    28: 'Subordinate clauses: որ, որը, ով, երբ, որտեղ',
    29: 'Everyday Yerevan speech',
    30: 'Punctuation and intonation marks: ։ ՝ ՞ ՛ ՜',
    31: 'Eastern and Western Armenian',
    32: 'Telling the time, dates and age',
    33: 'Word formation: −ություն, −ական, −ուհի, −իչ, −արան, −անոց',
}

# Шаблоны поиска примеров в словаре (регулярные выражения PostgreSQL, без учёта регистра).
# Границы слова — явно через «не армянская буква» (B … E), чтобы не зависеть от локали базы.
_B = '(^|[^ա-ևԱ-Ֆ])'
_E = '([^ա-ևԱ-Ֆ]|$)'
_AUX = '(եմ|ես|է|ենք|եք|են)'
_AUX_PAST = '(էի|էիր|էր|էինք|էիք|էին)'
GRAMMAR_PATTERNS = {
    1: r'(ու|և)',
    2: r'[թփքչցղռ]',
    3: _B + r'[ա-և]{2,}ներ(ը|ն|ի|ին|ից|ով|ում)?' + _E,     # -եր не берём: գիշեր «ночь», ընկեր «друг»
    4: _B + r'[ա-և]{2,}(ից|ով|ին)' + _E + '|' + _B
       + r'[ա-և]{2,}ում(?! ' + _AUX + '|' + ' ' + _AUX_PAST + ')' + _E,   # -ում, но не «գնում եմ» (настоящее время)
    5: _B + r'(իմ|քո|նրա|մեր|ձեր|նրանց) [ա-և]+' + _E,
    6: _B + r'(ավելի|ամենա[ա-և]+|քան)' + _E,
    7: _B + r'(մեկ|երկու|երեք|չորս|հինգ|վեց|յոթ|ութ|ինը|տասը|քսան|երեսուն|հարյուր|հազար|առաջին|[ա-և]+երորդ)' + _E,
    8: _B + r'(ես|դու|նա|մենք|դուք|նրանք|այս|այդ|այն|սա|դա)' + _E,
    9: _B + r'(ո՞վ|ի՞նչ|որտե՞ղ|ե՞րբ|ինչո՞ւ|ինչպե՞ս|ո՞ր|քանի՞|ու՞մ|ու՞ր)',
    10: r'[ա-և]{2,}ում ' + _AUX + _E,
    11: r'[ա-և]{2,}ում ' + _AUX_PAST + _E,
    12: _B + r'[ա-և]{2,}(եցի|եցիր|եց|եցինք|եցիք|եցին|ացի|ացիր|աց|ացինք|ացիք|ացին)' + _E,
    13: r'[ա-և]{2,}(ել|ացել) (' + _AUX[1:-1] + '|' + _AUX_PAST[1:-1] + ')' + _E,
    14: _B + r'(չ[ա-և]{2,}|մի՛)' + _E,
    15: r'[ա-և]՛[ա-և]|[ա-և]{2,}(իր|եք|եցեք|ացեք|ացրու)։$',
    16: r'^[^ ]+ [^ ]+ [^ ]+ ' + _AUX + '։$',
    17: _B + r'(կա|չկա|կան|չկան|ունեմ|ունես|ունի|ունենք|ունեք|ունեն|չունեմ|չունես|չունի)' + _E,
    18: _B + r'(և|բայց|կամ|որովհետև|ուստի|եթե|սակայն)' + _E,
    19: r'՞',
    20: r'«|ասաց, որ|ասում է, որ|հարցրեց',
    21: r'[ա-և]{2,}(ելու|ալու) ' + _AUX + _E + '|' + _B + r'կ[ա-և]{2,}(եմ|ենք|եք|ամ|անք|աք)' + _E,
    22: _B + r'(որ|որպեսզի|թող) [ա-և]{2,}(եմ|ես|ենք|եք|են|ամ|աս|անք|աք|ան)' + _E,
    23: _B + r'(եթե|Եթե)' + _E,
    24: _B + r'(պետք է|պետք չէ|կարող (եմ|ես|է|ենք|եք|են)|կարելի է|չի կարելի|պիտի)' + _E,
    25: _B + r'[ա-և]{3,}(ած|ելիս|ալիս|ելով|ալով)' + _E,
    26: _B + r'[ա-և]{2,}(վեց|վել|վում|ցնել|ցնում|ցրեց|ցրի)' + _E,
    27: _B + r'(վրա|մեջ|տակ|մոտ|հետ|համար|մասին|կողքին|պատճառով|փոխարեն)' + _E,
    28: r', (որ|որը|որի|որին|ով|երբ|որտեղ|որովհետև|քանի որ)' + _E,
    29: _B + r'(ա|էս|էդ|էն|կարամ|էգուց|ո՞նց|ինչքա՞ն|էստեղ)' + _E,
    30: r'[՞՛՜]',
    32: _B + r'(ժամը|տարեկան|թվականին|երկուշաբթի|երեքշաբթի|չորեքշաբթի|հինգշաբթի|ուրբաթ|կիրակի)' + _E,
    33: _B + r'[ա-և]{3,}(ություն|ական|ային|ուհի|արան|անոց)' + _E,
}


def clean(v):
    if v is None:
        return None
    s = str(v).strip()
    if not s or s in ('—', '-'):
        return None
    return s


def uz_apostrophes(s, uz=False):
    """В узбекской версии здесь приводились апострофы oʻ gʻ ʼ. Для армянского текста — только NFC и пробелы."""
    if s is None:
        return None
    import unicodedata
    return re.sub(r'[ \t]+', ' ', unicodedata.normalize('NFC', s))


def num(v):
    try:
        return int(float(v))
    except (TypeError, ValueError):
        return None


def q(v):
    """SQL-литерал."""
    if v is None:
        return 'null'
    if isinstance(v, int):
        return str(v)
    s = str(v)
    assert '$q$' not in s
    return f'$q${s}$q$'


def main_rows(path):
    ws = openpyxl.load_workbook(path, data_only=True).worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    header = [clean(h) for h in rows[3]]
    out = []
    for r in rows[4:]:
        if num(r[0]) is None:
            continue
        out.append(dict(zip(header, r)))
    return header, out


def col(row, *names):
    for n in names:
        if n in row:
            return clean(row[n])
    raise KeyError(names)


def build(src: Path):
    sql = ['-- Сгенерировано scripts/build_courses.py — не править вручную.', 'begin;']
    stats = {}

    for course, fname in SOURCES.items():
        header, rows = main_rows(src / fname)
        stats[course] = {'units': len(rows), 'tasks': 0}
        # разделы
        if course in SECTION_ORDER:
            for i, title in enumerate(SECTION_ORDER[course], 1):
                uz, en = SECTION_TR[title]
                sql.append(
                    f"insert into public.course_sections (course_id, title_ru, title_uz, title_en, sort_order) "
                    f"values ({q(course)}, {q(title)}, {q(uz)}, {q(en)}, {i}) "
                    f"on conflict (course_id, title_ru) do update set title_uz = excluded.title_uz, title_en = excluded.title_en, sort_order = excluded.sort_order;")
            known = set(SECTION_ORDER[course])
            for r in rows:
                sec = col(r, 'Раздел')
                if sec not in known:
                    raise SystemExit(f'{course}: неизвестный раздел «{sec}» — добавьте в SECTION_ORDER и SECTION_TR')

        for r in rows:
            n = num(r['№'])
            if course == 'grammar':
                full = col(r, 'Тема (RU / HY / EN)')
                m = re.match(r'^(.*?)\s*\(([^()]*[ա-ևԱ-Ֆ][^()]*(?:\([^()]*\)[^()]*)*)\)\s*$', full)
                title_ru, title_uz = (m.group(1).strip(), m.group(2).strip()) if m else (full, None)
                title_en = GRAMMAR_TITLES_EN.get(n)
                ex = ('Экспресс-объяснение (русский)', 'Экспресс-объяснение (հայերեն)', 'Express explanation (English)')
                de = ('Подробное объяснение (русский)', 'Подробное объяснение (հայերեն)', 'Detailed explanation (English)')
                tasks = [(f'Задание {k} (русский)', f'Задание {k} (հայերեն)', f'Task {k} (English)') for k in range(1, 6)]
                period = (None, None, None)
            else:
                title_ru = col(r, 'Тема (русский)')
                title_uz = col(r, 'Թեմա (հայերեն)')
                title_en = col(r, 'Topic (English)')
                ex = ('Экспресс-объяснение (русский)', 'Համառոտ բացատրություն (հայերեն)', 'Express explanation (English)')
                de = ('Подробное объяснение (русский)', 'Մանրամասն բացատրություն (հայերեն)', 'Detailed explanation (English)')
                tasks = [(f'Задание {k} (русский)', f'Առաջադրանք {k} (հայերեն)', f'Task {k} (English)') for k in range(1, 6)]
                period = (col(r, 'Период'), None, None) if course == 'history' else (None, None, None)

            section = col(r, 'Раздел') if 'Раздел' in r else None
            section_sql = (f"(select id from public.course_sections where course_id = {q(course)} and title_ru = {q(section)})"
                           if section else 'null')
            vals = {
                'title_ru': uz_apostrophes(title_ru), 'title_uz': uz_apostrophes(title_uz, True), 'title_en': uz_apostrophes(title_en),
                'express_ru': uz_apostrophes(clean(r[ex[0]])), 'express_uz': uz_apostrophes(clean(r[ex[1]]), True),
                'express_en': uz_apostrophes(clean(r[ex[2]])),
                'detailed_ru': uz_apostrophes(clean(r[de[0]])), 'detailed_uz': uz_apostrophes(clean(r[de[1]]), True),
                'detailed_en': uz_apostrophes(clean(r[de[2]])),
                'source': clean(r.get('Источник / материал')),
                'period': period[0],
                'example_pattern': GRAMMAR_PATTERNS.get(n) if course == 'grammar' else None,
                'cefr': UNIT_CEFR.get(course, {}).get(n, DEFAULT_CEFR[course]),
            }
            cols = ', '.join(vals)
            sql.append(
                f"insert into public.course_units (course_id, source_no, section_id, sort_order, {cols}) values "
                f"({q(course)}, {n}, {section_sql}, {n}, {', '.join(q(v) for v in vals.values())}) "
                f"on conflict (course_id, source_no) do update set section_id = excluded.section_id, sort_order = excluded.sort_order, "
                + ', '.join(f'{k} = excluded.{k}' for k in vals) + ';')
            unit_sql = f"(select id from public.course_units where course_id = {q(course)} and source_no = {n})"
            sql.append(f"delete from public.unit_terms where unit_id = {unit_sql};")
            for k in unit_terms(vals['express_ru'], vals['detailed_ru']):
                sql.append(f"insert into public.unit_terms (unit_id, term_key) values ({unit_sql}, {q(k)}) on conflict do nothing;")
            for k, (ru, uz, en) in enumerate(tasks, 1):
                pr, pu, pe = clean(r.get(ru)), clean(r.get(uz)), clean(r.get(en))
                if not pr:
                    continue
                stats[course]['tasks'] += 1
                sql.append(
                    f"insert into public.unit_tasks (unit_id, n, kind, prompt_ru, prompt_uz, prompt_en, origin, status) values "
                    f"({unit_sql}, {k}, 'open', {q(uz_apostrophes(pr))}, {q(uz_apostrophes(pu, True))}, {q(uz_apostrophes(pe))}, 'source', 'approved') "
                    f"on conflict (unit_id, n) do update set prompt_ru = excluded.prompt_ru, prompt_uz = excluded.prompt_uz, prompt_en = excluded.prompt_en "
                    f"where public.unit_tasks.origin = 'source';")
            # лишние темы, удалённые из файла, остаются в базе — удаляйте вручную, чтобы не потерять прогресс

    # ---- хронология истории
    wb = openpyxl.load_workbook(src / SOURCES['history'], data_only=True)
    ws = wb['Хронология (сводная таблица)']
    n_tl = 0
    for r in list(ws.iter_rows(values_only=True))[1:]:
        k = num(r[0])
        if k is None:
            continue
        n_tl += 1
        sql.append(
            f"insert into public.timeline_events (course_id, date_label, text_ru, text_uz, text_en, sort_order) values "
            f"('history', {q(clean(r[1]))}, {q(uz_apostrophes(clean(r[2])))}, {q(uz_apostrophes(clean(r[3]), True))}, {q(uz_apostrophes(clean(r[4])))}, {k}) "
            f"on conflict (course_id, sort_order) do update set date_label = excluded.date_label, text_ru = excluded.text_ru, "
            f"text_uz = excluded.text_uz, text_en = excluded.text_en;")
    stats['history']['timeline'] = n_tl

    # ---- справочник контактов (обществознание)
    ws = openpyxl.load_workbook(src / SOURCES['civics'], data_only=True)['Справочник контактов']
    n_c = 0
    for r in list(ws.iter_rows(values_only=True))[1:]:
        k = num(r[0])
        if k is None:
            continue
        n_c += 1
        sql.append(
            f"insert into public.reference_items (course_id, kind, title_ru, title_uz, title_en, value, sort_order) values "
            f"('civics', 'contact', {q(clean(r[1]))}, {q(uz_apostrophes(clean(r[2]), True))}, {q(clean(r[3]))}, {q(clean(r[4]))}, {k}) "
            f"on conflict (course_id, kind, sort_order) do update set title_ru = excluded.title_ru, title_uz = excluded.title_uz, "
            f"title_en = excluded.title_en, value = excluded.value;")
    stats['civics']['contacts'] = n_c

    # ---- базовая лексика (грамматика)
    ws = openpyxl.load_workbook(src / SOURCES['grammar'], data_only=True)['Базовая лексика (уроки 1–6)']
    n_v = 0
    for r in list(ws.iter_rows(values_only=True))[1:]:
        k = num(r[0])
        if k is None:
            continue
        n_v += 1
        extra = {'transcription': clean(r[2]), 'example': uz_apostrophes(clean(r[5]))}
        sql.append(
            f"insert into public.reference_items (course_id, kind, title_ru, title_uz, title_en, value, extra, sort_order) values "
            f"('grammar', 'vocab', {q(clean(r[3]))}, {q(uz_apostrophes(clean(r[1]), True))}, {q(clean(r[4]))}, {q(uz_apostrophes(clean(r[1]), True))}, "
            f"{q(json.dumps(extra, ensure_ascii=False))}::jsonb, {k}) "
            f"on conflict (course_id, kind, sort_order) do update set title_ru = excluded.title_ru, title_uz = excluded.title_uz, "
            f"title_en = excluded.title_en, value = excluded.value, extra = excluded.extra;")
    stats['grammar']['vocab'] = n_v

    # ---- вопросы для тестов по темам (черновики; публикует администратор в приложении)
    qpath = src.parent / 'quiz' / 'questions.json'
    n_q = 0
    if qpath.exists():
        for qq in json.loads(qpath.read_text(encoding='utf-8')):
            n_q += 1
            unit_sql = f"(select id from public.course_units where course_id = {q(qq['course'])} and source_no = {int(qq['n'])})"
            options = json.dumps({'ru': qq['options_ru'], 'uz': qq['options_uz'], 'en': qq['options_en']}, ensure_ascii=False)
            sql.append(
                f"insert into public.unit_tasks (unit_id, n, kind, prompt_ru, prompt_uz, prompt_en, options, answer, source_quote, origin, status) values "
                f"({unit_sql}, {100 + int(qq['k'])}, 'choice', {q(qq['q_ru'])}, {q(qq['q_uz'])}, {q(qq['q_en'])}, {q(options)}::jsonb, '0'::jsonb, "
                f"{q(qq['quote_ru'])}, 'generated', 'draft') "
                # изменился текст вопроса — снова на проверку; не изменился — статус сохраняется
                f"on conflict (unit_id, n) do update set "
                f"status = case when public.unit_tasks.prompt_ru is distinct from excluded.prompt_ru "
                f"or public.unit_tasks.options is distinct from excluded.options then 'draft' else public.unit_tasks.status end, "
                f"prompt_ru = excluded.prompt_ru, prompt_uz = excluded.prompt_uz, prompt_en = excluded.prompt_en, "
                f"options = excluded.options, answer = excluded.answer, source_quote = excluded.source_quote, kind = excluded.kind;")
    stats['quiz'] = {'questions': n_q}

    # ---- диалоги (data/source/dialogs.json)
    dpath = src / 'dialogs.json'
    if dpath.exists():
        dialogs = json.loads(dpath.read_text(encoding='utf-8'))
        stats['dialogs'] = {'units': len(dialogs), 'lines': 0}
        for i, title in enumerate(SECTION_ORDER['dialogs'], 1):
            uz, en = SECTION_TR[title]
            sql.append(
                f"insert into public.course_sections (course_id, title_ru, title_uz, title_en, sort_order) "
                f"values ('dialogs', {q(title)}, {q(uz)}, {q(en)}, {i}) "
                f"on conflict (course_id, title_ru) do update set title_uz = excluded.title_uz, title_en = excluded.title_en, sort_order = excluded.sort_order;")
        for d in dialogs:
            n = int(d['n'])
            if d['section'] not in SECTION_ORDER['dialogs']:
                raise SystemExit(f"dialogs.json: неизвестный раздел «{d['section']}»")
            stats['dialogs']['lines'] += len(d['lines'])
            extra = {'speakers': d['speakers'], 'lines': [{k: uz_apostrophes(v, True) if k == 'uz' else v for k, v in ln.items()} for ln in d['lines']],
                     'vocab': [uz_apostrophes(v, True) for v in d.get('vocab', [])]}
            vals = {
                'title_ru': d['title_ru'], 'title_uz': uz_apostrophes(d['title_uz'], True), 'title_en': d['title_en'],
                'express_ru': d['intro_ru'], 'express_uz': uz_apostrophes(d['intro_uz'], True), 'express_en': d['intro_en'],
                'detailed_ru': d.get('note_ru'), 'detailed_uz': uz_apostrophes(d.get('note_uz'), True), 'detailed_en': d.get('note_en'),
                'cefr': d['cefr'], 'extra': json.dumps(extra, ensure_ascii=False),
            }
            section_sql = f"(select id from public.course_sections where course_id = 'dialogs' and title_ru = {q(d['section'])})"
            sql.append(
                f"insert into public.course_units (course_id, source_no, section_id, sort_order, {', '.join(vals)}) values "
                f"('dialogs', {n}, {section_sql}, {n}, {', '.join(q(v) for v in vals.values())}) "
                f"on conflict (course_id, source_no) do update set section_id = excluded.section_id, sort_order = excluded.sort_order, "
                + ', '.join(f'{k} = excluded.{k}' for k in vals) + ';')
            unit_sql = f"(select id from public.course_units where course_id = 'dialogs' and source_no = {n})"
            sql.append(f"delete from public.unit_terms where unit_id = {unit_sql};")
            for v in d.get('vocab', []):
                k = key_uz(v)
                if k:
                    sql.append(f"insert into public.unit_terms (unit_id, term_key) values ({unit_sql}, {q(k)}) on conflict do nothing;")

    sql.append('select public.refresh_unit_words();')
    sql.append('commit;')
    return '\n'.join(sql) + '\n', stats


if __name__ == '__main__':
    src = Path(sys.argv[1] if len(sys.argv) > 1 else 'data/source')
    out = Path(sys.argv[2] if len(sys.argv) > 2 else 'data/build/courses.sql')
    out.parent.mkdir(parents=True, exist_ok=True)
    text, stats = build(src)
    out.write_text(text, encoding='utf-8')
    for c, s in stats.items():
        print(c, ', '.join(f'{k}: {v}' for k, v in s.items()))
    print(f'→ {out} ({len(text) // 1024} КБ)')
